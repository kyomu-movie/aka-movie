from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "member-guide-ja.md"
OUTPUT = ROOT / "docs" / "member-guide-ja.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
CALL_OUT = "F4F6F9"
CAUTION = "FFF4E5"
TABLE_HEADER = "E8EEF5"
MUTED = "5F6B76"
CONTENT_WIDTH_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_pr = table._tbl.tblPr
    width = table_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        table_pr.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = table_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    layout = table_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for index, grid_col in enumerate(grid.gridCol_lst):
        grid_col.set(qn("w:w"), str(widths[index]))
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None, code: bool = False) -> None:
    run.font.name = "Consolas" if code else "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas" if code else "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas" if code else "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_numbering(doc: Document, abstract_id: int, num_id: int, bullet: bool) -> None:
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def apply_list_number(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    code = doc.styles.add_style("Guide Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    code.font.size = Pt(9)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("aka-movie | メンバー向け手順書")
    set_run_font(header_run, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("aka-movie  |  ")
    set_run_font(footer_run, size=9, color=MUTED)
    add_page_field(footer)


def add_title(doc: Document, text: str) -> None:
    quick_guide = "これだけやればOK" in text
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8 if quick_guide else 36)
    paragraph.paragraph_format.space_after = Pt(4 if quick_guide else 8)
    run = paragraph.add_run(text)
    set_run_font(run, size=18 if quick_guide else 24, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8 if quick_guide else 22)
    run = subtitle.add_run("GitHub 参加から動画共有 PR 作成まで")
    set_run_font(run, size=10 if quick_guide else 12, color=MUTED)
    repo = doc.add_paragraph()
    repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repo.paragraph_format.space_after = Pt(10 if quick_guide else 26)
    add_hyperlink(repo, "https://github.com/kyomu-movie/aka-movie", "https://github.com/kyomu-movie/aka-movie")


def add_callout(doc: Document, kind: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CAUTION if kind == "CAUTION" else CALL_OUT)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label = paragraph.add_run("注意: " if kind == "CAUTION" else "重要: ")
    set_run_font(label, size=10.5, color="7A5A00" if kind == "CAUTION" else DARK_BLUE, bold=True)
    body = paragraph.add_run(text)
    set_run_font(body, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_markdown_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    if len(headers) == 2:
        widths = [2700, 6660]
    elif len(headers) == 3:
        widths = [2100, 3630, 3630]
    else:
        widths = [CONTENT_WIDTH_DXA // len(headers)] * len(headers)
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, TABLE_HEADER)
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=10, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, color=INK)
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow(doc: Document) -> None:
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(4)
    intro.paragraph_format.space_after = Pt(4)
    intro.paragraph_format.keep_with_next = True
    run = intro.add_run("動画作成・共有フロー")
    set_run_font(run, size=11, color=DARK_BLUE, bold=True)
    for index, item in enumerate((
        "台本を渡す",
        "画像・構造・日本語タイムラインを確認し、必要な段階で OK を返す",
        "MP4 を描画して確認し、最後の OK で共有する",
        "Codex が個人ブランチへ push・PR 作成し、オーナー確認を待つ",
        "オーナー承認後、Codex が学習案を提示し、承認された内容だけ反映する",
    ), start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.38)
        paragraph.paragraph_format.first_line_indent = Inches(-0.19)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = index < 5
        marker = paragraph.add_run(f"{index}. ")
        set_run_font(marker, size=10.5, color=BLUE, bold=True)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5, color=INK)
        if index < 5:
            arrow = doc.add_paragraph()
            arrow.paragraph_format.left_indent = Inches(0.38)
            arrow.paragraph_format.space_after = Pt(2)
            arrow.paragraph_format.keep_with_next = True
            run = arrow.add_run("↓")
            set_run_font(run, size=10, color=MUTED)


def render_markdown(doc: Document, source: str) -> None:
    lines = source.splitlines()
    index = 0
    next_abstract = 91
    next_num = 91
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            add_title(doc, line[2:].strip())
            index += 1
            continue
        if line.startswith("## "):
            if line.startswith("## Codex "):
                doc.add_page_break()
            paragraph = doc.add_paragraph(line[3:].strip(), style="Heading 1")
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(line[4:].strip(), style="Heading 2")
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if line.startswith("> [!"):
            kind = re.search(r"\[!(\w+)\]", line).group(1)
            body = lines[index + 1].lstrip("> ").strip() if index + 1 < len(lines) else ""
            add_callout(doc, kind, body)
            index += 2
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            if language == "mermaid":
                add_flow(doc)
            else:
                paragraph = doc.add_paragraph(style="Guide Code")
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.right_indent = Inches(0.18)
                paragraph.paragraph_format.keep_together = True
                run = paragraph.add_run("\n".join(block))
                set_run_font(run, size=9, color=INK, code=True)
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s|:-]+\|$", lines[index + 1]):
            headers = parse_table_row(line)
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(parse_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, headers, rows)
            continue
        if re.match(r"^\d+\. ", line):
            entries: list[str] = []
            while index < len(lines) and re.match(r"^\d+\. ", lines[index]):
                entries.append(re.sub(r"^\d+\. ", "", lines[index]))
                index += 1
            add_numbering(doc, next_abstract, next_num, bullet=False)
            for item_index, entry in enumerate(entries):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = item_index < len(entries) - 1
                apply_list_number(paragraph, next_num)
                run = paragraph.add_run(entry)
                set_run_font(run, size=11, color=INK)
            next_abstract += 1
            next_num += 1
            continue
        if line.startswith("- ") and not line.startswith("- [ ] "):
            entries: list[str] = []
            while index < len(lines) and lines[index].startswith("- "):
                entries.append(lines[index][2:])
                index += 1
            add_numbering(doc, next_abstract, next_num, bullet=True)
            for item_index, entry in enumerate(entries):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = item_index < len(entries) - 1
                apply_list_number(paragraph, next_num)
                run = paragraph.add_run(entry)
                set_run_font(run, size=11, color=INK)
            next_abstract += 1
            next_num += 1
            continue
        if line.startswith("- [ ] "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.375)
            paragraph.paragraph_format.first_line_indent = Inches(-0.19)
            paragraph.paragraph_format.space_after = Pt(2)
            checkbox = paragraph.add_run("□ ")
            set_run_font(checkbox, size=11, color=BLUE, bold=True)
            run = paragraph.add_run(line[6:])
            set_run_font(run, size=11, color=INK)
            index += 1
            continue
        if line.startswith("    "):
            paragraph = doc.add_paragraph(style="Guide Code")
            run = paragraph.add_run(line.strip())
            set_run_font(run, size=9, color=INK, code=True)
            index += 1
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(line)
        set_run_font(run, size=11, color=INK)
        index += 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a member guide DOCX from Markdown.")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    doc = Document()
    configure_document(doc)
    render_markdown(doc, args.source.read_text(encoding="utf-8"))
    doc.save(args.output)
    print(f"Created {args.output}")


if __name__ == "__main__":
    main()
